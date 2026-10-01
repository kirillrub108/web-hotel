"""Проверки собранного отчёта. Запускается из build_report.py; при ошибке — SystemExit со списком проблем."""
import json
import posixpath
import re
import sys
import zipfile
from pathlib import Path

import docx
from lxml import etree

from title import q, text_of

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

TITLE_MUST = ["Рубец Кирилл Сергеевич", "Рубцу Кириллу Сергеевичу", "521428", "Разработка web-приложения гостиницы"]
FORBIDDEN = ["521328", "НАЗВАНИЕ ВАШЕЙ ТЕМЫ", "КОНТРОЛЬНАЯ РАБОТА", "Здесь должно быть", "ЗДЕСЬ ДОЛЖНО БЫТЬ",
             "НАЗВАНИЕ ОРГАНИЗАЦИИ", "НАЗВАНИЕ СРЕДСТВА", "ТЕМА КУРСОВОЙ РАБОТЫ", "Аналог 1", "Аналог 2",
             "Название аналога", "И т.д.", "……", "…..", "Вариант использования 1", "мастер-класс", "Мастер-класс",
             "мастерск", "Strapi", "Тярина", "Django", "settings.py", "views.py", "urls.py", "Замечание!!!",
             "ДАЛЕЕ ПРИВЕДЕН ПРИМЕР", "(-аяся)", "Выполнил (-а)", "Html- код"]
FORBIDDEN_CASE = [r"\bРИСУНОК\b", r"\bФИО\.", r"«ТЕМА»"]


def _parts(z: zipfile.ZipFile) -> dict[str, bytes]:
    return {n: z.read(n) for n in z.namelist()}


def _rels(parts: dict, part: str) -> dict[str, str]:
    folder, name = part.rsplit("/", 1)
    rel_path = f"{folder}/_rels/{name}.rels"
    if rel_path not in parts:
        return {}
    root = etree.fromstring(parts[rel_path])
    return {r.get("Id"): r.get("Target") for r in root}


def check_package(path: Path, problems: list) -> str:
    docx.Document(str(path))  # открывается python-docx
    z = zipfile.ZipFile(path)
    parts = _parts(z)
    used_media = set()
    for part in [p for p in parts if p.startswith("word/") and p.endswith(".xml") and "/_rels/" not in p]:
        rels = _rels(parts, part)
        xml = parts[part].decode("utf-8")
        for rid in set(re.findall(r'r:(?:id|embed|link)="([^"]+)"', xml)):
            if rid not in rels:
                problems.append(f"{part}: битая ссылка {rid}")
            elif rels[rid].startswith("media/"):
                used_media.add("word/" + rels[rid])
    doc_xml = parts["word/document.xml"].decode("utf-8")
    numbering = parts["word/numbering.xml"].decode("utf-8")
    nums = set(re.findall(r'<w:num w:numId="(\d+)"', numbering)) | {"0"}
    for n in set(re.findall(r'<w:numId w:val="(\d+)"/>', doc_xml)) - nums:
        problems.append(f"numId {n} нет в numbering.xml")
    for target in _rels(parts, "word/document.xml").values():
        if not target.startswith("http") and posixpath.normpath(f"word/{target}") not in parts:
            problems.append(f"связь указывает на отсутствующий файл {target}")
    ours = {m for m in parts if m.startswith("word/media/report_")}
    if ours - used_media:
        problems.append(f"медиа не вставлены в документ: {sorted(ours - used_media)}")
    ids = re.findall(r'<wp:docPr id="(\d+)"', doc_xml)
    if len(ids) != len(set(ids)):
        problems.append("повторяющиеся docPr id")
    if 'w:fldCharType="end"' not in doc_xml or "updateFields" not in parts["word/settings.xml"].decode():
        problems.append("поле TOC без end или нет updateFields")
    return doc_xml


def paragraphs(doc_xml: str) -> list[tuple[str, str, str]]:
    """(styleId, numId из pPr, текст) для всех абзацев тела."""
    body = etree.fromstring(doc_xml.encode()).find(q("body"))
    out = []
    for p in body.iter(q("p")):
        st = p.find(f"{q('pPr')}/{q('pStyle')}")
        num = p.find(f"{q('pPr')}/{q('numPr')}/{q('numId')}")
        out.append((st.get(q("val")) if st is not None else "", num.get(q("val")) if num is not None else "",
                    text_of(p)))
    return out


def old_kr_paragraphs() -> set[str]:
    """Длинные абзацы тела старой контрольной из «Шаблона КР» — ни один не должен попасть в отчёт."""
    z = zipfile.ZipFile(ROOT / "docs" / "templates" / "Шаблон КР.docx")
    body = etree.fromstring(z.read("word/document.xml")).find(q("body"))
    styles = {s.get(q("styleId")): s.find(q("name")).get(q("val"))
              for s in etree.fromstring(z.read("word/styles.xml")).iter(q("style"))}
    texts = []
    for p in body.iter(q("p")):
        st = p.find(f"{q('pPr')}/{q('pStyle')}")
        texts.append((styles.get(st.get(q("val"))) if st is not None else "", text_of(p).strip()))
    start = next(i for i, (_, t) in enumerate(texts) if t.upper() == "ОГЛАВЛЕНИЕ")
    # Строки кода старых листингов (стиль «Листинг») не сравниваются: import-строки совпадают у любых проектов.
    return {t for s, t in texts[start + 1:] if len(t) > 40 and s != "Листинг"}


def compose_secret_values() -> list[str]:
    compose = (ROOT / "docker-compose.yml").read_text(encoding="utf-8")
    return [v for k, v in re.findall(r"\$\{(\w*(?:PASSWORD|SECRET|TOKEN|KEY)\w*):-([^}]*)\}", compose)
            if len(v) >= 8 and "MIN_LENGTH" not in k]


def run(path: Path, nums: dict, items: list) -> None:
    problems: list[str] = []
    doc_xml = check_package(path, problems)
    paras = paragraphs(doc_xml)
    full = "\n".join(t for _, _, t in paras)

    # Титул и лист задания — всё до ОГЛАВЛЕНИЯ.
    toc_at = next(i for i, (_, _, t) in enumerate(paras) if t.strip().upper() == "ОГЛАВЛЕНИЕ")
    title = "\n".join(t for _, _, t in paras[:toc_at])
    problems += [f"на титуле нет «{s}»" for s in TITLE_MUST if s not in title]

    smap = json.loads((HERE / "style_map.json").read_text(encoding="utf-8"))
    problems += [f"остался маркер шаблона/старой работы «{s}»" for s in FORBIDDEN if s in full]
    problems += [f"остался маркер {rx}" for rx in FORBIDDEN_CASE if re.search(rx, full)]
    prose = "\n".join(t for s, _, t in paras if s != smap["listing_body"])
    leaked = [t[:60] for t in old_kr_paragraphs() if t in prose]
    problems += [f"абзац старой контрольной: «{t}…»" for t in leaked]

    # Подписи: число автонумеруемых подписей = числу рисунков/таблиц в контенте.
    figs = [t for s, n, t in paras if s == smap["figure_caption"]]
    tabs = [t for s, n, t in paras if s == smap["table_caption"]]
    lst_main = [t for s, n, t in paras if s == smap["listing_caption"] and n != "0"]
    if len(figs) != len(nums["fig"]):
        problems.append(f"подписей рисунков {len(figs)}, а рисунков в контенте {len(nums['fig'])}")
    if len(tabs) != len(nums["tab"]):
        problems.append(f"подписей таблиц {len(tabs)}, а таблиц в контенте {len(nums['tab'])}")
    if lst_main:
        problems.append(f"в основной части есть автонумеруемые листинги: {lst_main}")
    for n in range(1, len(nums["ref"]) + 1):
        if f"[{n}]" not in full:
            problems.append(f"на источник {n} нет ссылки [{n}]")

    placeholders = re.findall(r"\[Место для рисунка: ([^\]]+)\]", full)
    shots = [it for it in items if it[0] == "fig" and "/screens/" in "/" + it[2] and not it[2].startswith("screens/analog")]
    for it in shots:
        if not (HERE.parent / "assets" / it[2]).exists():
            problems.append(f"нет скриншота {it[2]}")

    # Секреты в приложениях.
    appendix = full[full.find("ПРИЛОЖЕНИЕ А"):]
    for value in compose_secret_values():
        if value in full:
            problems.append("в документе значение секрета из docker-compose.yml")
    if re.search(r"\b[a-z][a-z0-9+.-]*://[^\s'\"/@]*:[^\s'\"@*]+@", appendix):
        problems.append("в приложениях строка подключения с паролем")
    allowed = ("***", "os.getenv", "int(os.getenv", "${", "{")
    for line in appendix.splitlines():
        m = re.match(r"\s*-?\s*([A-Z0-9_]*(?:PASSWORD|SECRET|TOKEN|API_KEY|DATABASE_URL)[A-Z0-9_]*)\s*[:=]\s*(.*)",
                     line)
        if not m or m.group(1).endswith(("_MIN_LENGTH", "_HASH", "_TTL")):
            continue
        value = m.group(2).strip()
        if value and not value.startswith(allowed) and value.strip("\"'") != m.group(1).lower():
            problems.append(f"возможный секрет в приложении: {line.strip()[:80]}")

    # Сокращения: каждое встречается в тексте вне своего раздела.
    from content.front import ABBREVIATIONS
    start = full.find("Определения, обозначения и сокращения")
    end = full.find("Введение", start)
    rest = full[:start] + full[end:]
    for abbr, _ in ABBREVIATIONS:
        if not re.search(rf"(?<![\w-]){re.escape(abbr)}(?![\w])", rest):
            problems.append(f"сокращение {abbr} не встречается в тексте")

    print(f"[verify] рисунков: {len(figs)}, таблиц: {len(tabs)}, источников: {len(nums['ref'])}, "
          f"скриншотов сайта: {len(shots)}, заглушек рисунков: {len(placeholders)} {placeholders}")
    if problems:
        print("[verify] ПРОБЛЕМЫ:\n  " + "\n  ".join(problems), file=sys.stderr)
        raise SystemExit(1)
    print("[verify] OK")
